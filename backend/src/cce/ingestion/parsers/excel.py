
import openpyxl
from typing import Union, BinaryIO
from cce.ingestion.parsers.base import DocumentParser
from cce.ingestion.models import (ProcessingResult, DocumentMetadata, ProcessingStatus, 
                              CanonicalDocument, SpreadsheetElement, TableElement, TableCell)

class ExcelParser(DocumentParser):
    def parse(self, file_stream_or_path: Union[str, BinaryIO], metadata: DocumentMetadata) -> ProcessingResult:
        try:
            wb = openpyxl.load_workbook(file_stream_or_path, data_only=True)
            elements = []
            
            for sheet_idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                cells = []
                
                for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                    for col_idx, cell_value in enumerate(row):
                        if cell_value is not None:
                            cells.append(TableCell(
                                row=row_idx,
                                col=col_idx,
                                text=str(cell_value),
                                is_header=(row_idx == 0)
                            ))
                
                table = TableElement(
                    id=f"{metadata.document_id}_{sheet_name}_table",
                    rows=sheet.max_row,
                    cols=sheet.max_column,
                    cells=cells,
                    order=0
                )
                
                spreadsheet_elem = SpreadsheetElement(
                    id=f"{metadata.document_id}_{sheet_name}",
                    sheet_name=sheet_name,
                    order=sheet_idx,
                    tables=[table]
                )
                elements.append(spreadsheet_elem)
                
            doc = CanonicalDocument(metadata=metadata, elements=elements)
            return ProcessingResult(status=ProcessingStatus.SUCCESS, document=doc)
        except Exception as e:
            return ProcessingResult(status=ProcessingStatus.FAILED, errors=[str(e)])
