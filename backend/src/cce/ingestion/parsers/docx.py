import re
from importlib import metadata as package_metadata
from typing import BinaryIO, Union

from docx import Document

from cce.ingestion.models import (
    CanonicalDocument,
    DocumentMetadata,
    ElementType,
    ProcessingResult,
    ProcessingStatus,
    TableCell,
    TableElement,
    TextElement,
)
from cce.ingestion.parsers.base import DocumentParser


class DocxParser(DocumentParser):
    def parse(
        self,
        file_stream_or_path: Union[str, BinaryIO],
        metadata: DocumentMetadata,
    ) -> ProcessingResult:
        metadata.parser_used = "python-docx"
        metadata.parser_version = _package_version("python-docx")

        try:
            document = Document(file_stream_or_path)
            elements = []
            order = 0
            positions = {node: i for i, node in enumerate(document.element.body)}
            element_positions = {}

            for paragraph_index, paragraph in enumerate(document.paragraphs):
                text = paragraph.text.strip()
                if not text:
                    continue

                style_name = paragraph.style.name if paragraph.style else ""
                element_positions[f"{metadata.document_id}_elem_{order}"] = positions[paragraph._p]
                elements.append(
                    TextElement(
                        id=f"{metadata.document_id}_elem_{order}",
                        type=_paragraph_type(style_name),
                        text=text,
                        order=order,
                        page_number=None,
                        metadata={
                            "paragraph_index": paragraph_index,
                            "style": style_name,
                            "heading_level": int(re.search(r"\d+", style_name).group()) if re.search(r"\d+", style_name) and style_name.lower().startswith("heading") else None,
                        },
                    )
                )
                order += 1

            for table_index, table in enumerate(document.tables):
                rows = len(table.rows)
                cols = max((len(row.cells) for row in table.rows), default=0)
                cells = []
                for row_index, row in enumerate(table.rows):
                    for col_index, cell in enumerate(row.cells):
                        cells.append(
                            TableCell(
                                row=row_index,
                                col=col_index,
                                text=_cell_text(cell),
                                is_header=row_index == 0,
                            )
                        )

                if rows and cols:
                    element_positions[f"{metadata.document_id}_elem_{order}"] = positions[table._tbl]
                    elements.append(
                        TableElement(
                            id=f"{metadata.document_id}_elem_{order}",
                            rows=rows,
                            cols=cols,
                            cells=cells,
                            order=order,
                            page_number=None,
                            metadata={"table_index": table_index},
                        )
                    )
                    order += 1

            elements.sort(key=lambda element: element_positions[element.id])
            for position, element in enumerate(elements):
                element.order = position

            if not elements:
                return ProcessingResult(
                    status=ProcessingStatus.UNSUPPORTED,
                    errors=["No extractable DOCX text or tables found."],
                )

            return ProcessingResult(
                status=ProcessingStatus.SUCCESS,
                document=CanonicalDocument(metadata=metadata, elements=elements),
            )
        except Exception as exc:
            return ProcessingResult(
                status=ProcessingStatus.FAILED,
                errors=[str(exc)],
            )


def _paragraph_type(style_name: str) -> ElementType:
    normalized = style_name.lower()
    if normalized.startswith("heading"):
        return ElementType.HEADING
    if normalized.startswith("list") or " list" in normalized:
        return ElementType.LIST_ITEM
    return ElementType.PARAGRAPH


def _cell_text(cell) -> str:
    return "\n".join(
        paragraph.text.strip()
        for paragraph in cell.paragraphs
        if paragraph.text.strip()
    )


def _package_version(package_name: str) -> str:
    try:
        return package_metadata.version(package_name)
    except package_metadata.PackageNotFoundError:
        return "unknown"
