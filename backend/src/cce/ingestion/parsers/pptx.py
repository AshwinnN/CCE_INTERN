from importlib import metadata as package_metadata
from typing import BinaryIO, Union

from pptx import Presentation

from cce.ingestion.models import (
    CanonicalDocument,
    DocumentMetadata,
    ElementType,
    ProcessingResult,
    ProcessingStatus,
    TableCell,
    TableElement,
    TextElement,
)
from cce.ingestion.parsers.base import DocumentParser


class PptxParser(DocumentParser):
    def parse(
        self,
        file_stream_or_path: Union[str, BinaryIO],
        metadata: DocumentMetadata,
    ) -> ProcessingResult:
        metadata.parser_used = "python-pptx"
        metadata.parser_version = _package_version("python-pptx")

        try:
            presentation = Presentation(file_stream_or_path)
            elements = []
            order = 0

            for slide_index, slide in enumerate(presentation.slides):
                slide_number = slide_index + 1
                for shape_index, shape in enumerate(slide.shapes):
                    if getattr(shape, "has_text_frame", False):
                        for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                            text = paragraph.text.strip()
                            if not text:
                                continue
                            elements.append(
                                TextElement(
                                    id=f"{metadata.document_id}_elem_{order}",
                                    type=_paragraph_type(paragraph),
                                    text=text,
                                    order=order,
                                    page_number=slide_number,
                                    metadata={
                                        "slide_number": slide_number,
                                        "shape_index": shape_index,
                                        "paragraph_index": paragraph_index,
                                    },
                                )
                            )
                            order += 1

                    if getattr(shape, "has_table", False):
                        table = shape.table
                        rows = len(table.rows)
                        cols = len(table.columns)
                        cells = []
                        for row_index, row in enumerate(table.rows):
                            for col_index, cell in enumerate(row.cells):
                                cells.append(
                                    TableCell(
                                        row=row_index,
                                        col=col_index,
                                        text=cell.text.strip(),
                                        is_header=row_index == 0,
                                    )
                                )

                        if rows and cols:
                            elements.append(
                                TableElement(
                                    id=f"{metadata.document_id}_elem_{order}",
                                    rows=rows,
                                    cols=cols,
                                    cells=cells,
                                    order=order,
                                    page_number=slide_number,
                                    metadata={
                                        "slide_number": slide_number,
                                        "shape_index": shape_index,
                                    },
                                )
                            )
                            order += 1

            if not elements:
                return ProcessingResult(
                    status=ProcessingStatus.UNSUPPORTED,
                    errors=["No extractable PPTX text or tables found."],
                )

            return ProcessingResult(
                status=ProcessingStatus.SUCCESS,
                document=CanonicalDocument(metadata=metadata, elements=elements),
            )
        except Exception as exc:
            return ProcessingResult(
                status=ProcessingStatus.FAILED,
                errors=[str(exc)],
            )


def _paragraph_type(paragraph) -> ElementType:
    if getattr(paragraph, "level", 0) > 0:
        return ElementType.LIST_ITEM
    return ElementType.PARAGRAPH


def _package_version(package_name: str) -> str:
    try:
        return package_metadata.version(package_name)
    except package_metadata.PackageNotFoundError:
        return "unknown"
