import importlib
import sys

from cce.ingestion.change_detection.file_detection import detect_mime_type
from cce.ingestion.models import DocumentMetadata, ElementType, ProcessingStatus
from cce.ingestion.parsers.docx import DocxParser
from cce.ingestion.parsers.factory import ParserFactory
from cce.ingestion.parsers.pdf_text import PdfTextParser
from cce.ingestion.parsers.pptx import PptxParser


PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _metadata(document_id="fixture"):
    return DocumentMetadata(
        document_id=document_id,
        source_system="unit-test",
        source_uri=f"unit-test://{document_id}",
    )


def _write_docx_fixture(path):
    from docx import Document

    document = Document()
    document.add_paragraph("Contract Summary", style="Heading 1")
    document.add_paragraph("Customer ABC has parser coverage.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "SLA"
    table.cell(1, 1).text = "80%"
    document.save(path)


def _write_empty_docx(path):
    from docx import Document

    Document().save(path)


def _write_pptx_fixture(path):
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Parser Slide"
    textbox = slide.shapes.add_textbox(100000, 1000000, 5000000, 500000)
    textbox.text_frame.text = "PowerPoint parser coverage."
    table_shape = slide.shapes.add_table(2, 2, 100000, 1700000, 4000000, 800000)
    table = table_shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Coverage"
    table.cell(1, 1).text = "Yes"
    presentation.save(path)


def _write_empty_pptx(path):
    from pptx import Presentation

    Presentation().save(path)


def _write_pdf_fixture(path, text="Synthetic PDF parser fixture"):
    content = f"BT /F1 24 Tf 100 700 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
    ]
    _write_pdf(path, objects)


def _write_empty_pdf(path):
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length 0 >>\nstream\n\nendstream",
    ]
    _write_pdf(path, objects)


def _write_pdf(path, objects):
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(b"%d 0 obj\n" % object_number)
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(b"xref\n0 %d\n" % (len(objects) + 1))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(b"%010d 00000 n \n" % offset)
    pdf.extend(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref_offset)
    )
    path.write_bytes(bytes(pdf))


def test_factory_resolves_pdf_docx_and_pptx_to_split_parsers():
    assert isinstance(ParserFactory.get_parser(PDF_MIME), PdfTextParser)
    assert isinstance(ParserFactory.get_parser(DOCX_MIME), DocxParser)
    assert isinstance(ParserFactory.get_parser(PPTX_MIME), PptxParser)


def test_missing_dependency_for_one_parser_does_not_disable_others(monkeypatch):
    original_registry = dict(ParserFactory._registry)
    ParserFactory._registry = {}
    real_import_module = importlib.import_module

    def fake_import_module(module_path, package=None):
        if module_path == "cce.ingestion.parsers.docx":
            raise ImportError("missing python-docx")
        return real_import_module(module_path, package)

    monkeypatch.setattr(importlib, "import_module", fake_import_module)
    try:
        ParserFactory._try_register([DOCX_MIME], "cce.ingestion.parsers.docx", "DocxParser")
        ParserFactory._try_register([PDF_MIME], "cce.ingestion.parsers.pdf_text", "PdfTextParser")
        assert ParserFactory.get_parser(DOCX_MIME) is None
        assert isinstance(ParserFactory.get_parser(PDF_MIME), PdfTextParser)
    finally:
        ParserFactory._registry = original_registry


def test_mime_detection_falls_back_to_extension_for_generic_magic(monkeypatch, tmp_path):
    class FakeMagic:
        @staticmethod
        def from_file(file_path, mime=True):
            return "application/octet-stream"

    monkeypatch.setitem(sys.modules, "magic", FakeMagic)
    docx_path = tmp_path / "fixture.docx"
    docx_path.write_bytes(b"not parsed here")

    assert detect_mime_type(str(docx_path)) == DOCX_MIME


def test_pdf_parser_extracts_text(tmp_path):
    pdf_path = tmp_path / "fixture.pdf"
    _write_pdf_fixture(pdf_path)

    result = PdfTextParser().parse(str(pdf_path), _metadata("fixture.pdf"))

    assert result.status == ProcessingStatus.SUCCESS
    assert result.document.metadata.parser_used == "pypdfium2"
    assert result.document.elements[0].page_number == 1
    assert result.document.elements[0].text == "Synthetic PDF parser fixture"


def test_pdf_parser_empty_is_unsupported(tmp_path):
    pdf_path = tmp_path / "empty.pdf"
    _write_empty_pdf(pdf_path)

    result = PdfTextParser().parse(str(pdf_path), _metadata("empty.pdf"))

    assert result.status == ProcessingStatus.UNSUPPORTED
    assert result.document is None


def test_pdf_parser_malformed_is_failed(tmp_path):
    pdf_path = tmp_path / "bad.pdf"
    pdf_path.write_bytes(b"not a pdf")

    result = PdfTextParser().parse(str(pdf_path), _metadata("bad.pdf"))

    assert result.status == ProcessingStatus.FAILED
    assert result.errors


def test_docx_parser_extracts_text_and_table(tmp_path):
    docx_path = tmp_path / "fixture.docx"
    _write_docx_fixture(docx_path)

    result = DocxParser().parse(str(docx_path), _metadata("fixture.docx"))

    assert result.status == ProcessingStatus.SUCCESS
    assert result.document.metadata.parser_used == "python-docx"
    assert result.document.elements[0].type == ElementType.HEADING
    assert result.document.elements[1].type == ElementType.PARAGRAPH
    table = result.document.elements[2]
    assert table.rows == 2
    assert table.cols == 2
    assert table.cells[0].is_header is True
    assert table.cells[3].text == "80%"


def test_docx_parser_empty_is_unsupported(tmp_path):
    docx_path = tmp_path / "empty.docx"
    _write_empty_docx(docx_path)

    result = DocxParser().parse(str(docx_path), _metadata("empty.docx"))

    assert result.status == ProcessingStatus.UNSUPPORTED
    assert result.document is None


def test_docx_parser_malformed_is_failed(tmp_path):
    docx_path = tmp_path / "bad.docx"
    docx_path.write_bytes(b"not a docx")

    result = DocxParser().parse(str(docx_path), _metadata("bad.docx"))

    assert result.status == ProcessingStatus.FAILED
    assert result.errors


def test_pptx_parser_extracts_text_and_table(tmp_path):
    pptx_path = tmp_path / "fixture.pptx"
    _write_pptx_fixture(pptx_path)

    result = PptxParser().parse(str(pptx_path), _metadata("fixture.pptx"))

    assert result.status == ProcessingStatus.SUCCESS
    assert result.document.metadata.parser_used == "python-pptx"
    assert result.document.elements[0].page_number == 1
    assert result.document.elements[0].metadata["slide_number"] == 1
    assert any(
        element.type == ElementType.TABLE and element.cells[3].text == "Yes"
        for element in result.document.elements
    )


def test_pptx_parser_empty_is_unsupported(tmp_path):
    pptx_path = tmp_path / "empty.pptx"
    _write_empty_pptx(pptx_path)

    result = PptxParser().parse(str(pptx_path), _metadata("empty.pptx"))

    assert result.status == ProcessingStatus.UNSUPPORTED
    assert result.document is None


def test_pptx_parser_malformed_is_failed(tmp_path):
    pptx_path = tmp_path / "bad.pptx"
    pptx_path.write_bytes(b"not a pptx")

    result = PptxParser().parse(str(pptx_path), _metadata("bad.pptx"))

    assert result.status == ProcessingStatus.FAILED
    assert result.errors
