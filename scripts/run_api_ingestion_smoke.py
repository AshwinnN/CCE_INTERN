#!/usr/bin/env python3
"""API-only local-fs ingestion smoke test over synthetic content."""

from __future__ import annotations

import json
import os
import sys
import time
from pathlib import Path
from urllib import request

from cce.integrations.agentic_plane.local_index import LocalIndexClient


PHRASES = {
    "synthetic.txt": "Synthetic local-fs document with searchable plumbing text.",
    "synthetic.pdf": "Synthetic PDF parser coverage phrase.",
    "synthetic.docx": "Synthetic DOCX parser coverage phrase.",
    "synthetic.pptx": "Synthetic PPTX parser coverage phrase.",
}


def _json(method: str, url: str, payload: dict | None = None) -> dict:
    data = json.dumps(payload).encode("utf-8") if payload is not None else None
    req = request.Request(
        url,
        data=data,
        method=method,
        headers={"Content-Type": "application/json"},
    )
    with request.urlopen(req, timeout=20) as response:
        return json.loads(response.read().decode("utf-8"))


def main() -> int:
    base_url = os.environ.get("CCE_HTTP_BASE_URL", "http://localhost:8080")
    root = Path(os.environ.get("CCE_API_SMOKE_ROOT", "/tmp/cce-api-ingestion-smoke"))
    root.mkdir(parents=True, exist_ok=True)
    _write_corpus(root)

    registered = _json(
        "POST",
        f"{base_url}/sources",
        {
            "adapter": "local-fs",
            "source_id": "api-smoke-local-fs",
            "kind": "unstructured",
            "credential_ref": "",
            "config": {"root_path": str(root)},
        },
    )
    source_id = registered["source_id"]
    print(f"Registered source: {source_id}")

    connection = _json("POST", f"{base_url}/sources/{source_id}/test", {})
    if connection["status"] != "CONNECTED":
        raise RuntimeError(f"connection failed: {connection}")
    print("Connection: CONNECTED")

    triggered = _json("POST", f"{base_url}/sources/{source_id}/ingest", {})
    run_id = triggered["ingestion_run_id"]
    print(f"Triggered run: {run_id}")

    deadline = time.time() + 60
    status = triggered
    while time.time() < deadline and status["status"] == "RUNNING":
        time.sleep(1)
        status = _json("GET", f"{base_url}/ingestion-runs/{run_id}")
    if status["status"] != "SUCCESS":
        raise RuntimeError(f"ingestion did not succeed: {status}")
    print(f"Ingestion: {status['objects_processed']} processed")

    dsn = os.environ["CCE_CONTROL_DATABASE_URL"]
    index = LocalIndexClient(dsn)
    for filename, phrase in PHRASES.items():
        results = index.search(phrase, limit=5)
        if not results:
            raise RuntimeError(f"index search returned no results for {filename}")
        if results[0]["document_id"] != filename:
            raise RuntimeError(
                f"top search hit for {filename} was {results[0]['document_id']}: "
                f"{results[0]['provenance']}"
            )
        print(f"Search hit for {filename}:", results[0]["provenance"])
    return 0


def _write_corpus(root: Path) -> None:
    (root / "synthetic.txt").write_text(PHRASES["synthetic.txt"], encoding="utf-8")
    _write_pdf(root / "synthetic.pdf", PHRASES["synthetic.pdf"])
    _write_docx(root / "synthetic.docx", PHRASES["synthetic.docx"])
    _write_pptx(root / "synthetic.pptx", PHRASES["synthetic.pptx"])


def _write_docx(path: Path, phrase: str) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph(phrase)
    document.save(path)


def _write_pptx(path: Path, phrase: str) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = phrase
    presentation.save(path)


def _write_pdf(path: Path, phrase: str) -> None:
    content = f"BT /F1 24 Tf 100 700 Td ({phrase}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(b"%d 0 obj\n" % object_number)
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(b"xref\n0 %d\n" % (len(objects) + 1))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(b"%010d 00000 n \n" % offset)
    pdf.extend(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref_offset)
    )
    path.write_bytes(bytes(pdf))


if __name__ == "__main__":
    sys.exit(main())
